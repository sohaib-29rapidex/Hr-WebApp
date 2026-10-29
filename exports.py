"""Excel exports. The monthly sheet copies the RapidEx HR sheet layout:
title bar, NO./Office/Department/Emp Name/Total Absence, one column per day with a code dropdown,
colours that follow the code (so HR can still edit it in Excel), and live COUNTIF totals."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import engine

PLUM, ORANGE = "3B0D3F", "F26522"
CODE_STYLE = {"P": ("DCF3D4", "23701A"), "L": ("ECE2FB", "6A3AA3"), "A": ("FDDCDC", "C62828"),
              "H": ("DCE8FF", "1F4FC4"), "E": ("FFE7C4", "985200"), "S": ("3B0D3F", "FFFFFF"),
              "WH": ("1F3A8A", "FFFFFF"), "WO": ("EFECF0", "6B6070")}
thin = Side(style="thin", color="E6E0E8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fill(c):
    return PatternFill("solid", fgColor=c)


def _head(ws, row, values, fill=PLUM, color="FFFFFF"):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font, c.alignment, c.border = _fill(fill), Font(bold=True, color=color), CENTER, BOX


def _rows(ws, start, rows):
    for r_i, row in enumerate(rows, start):
        for c_i, v in enumerate(row, 1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.border = BOX
            c.alignment = Alignment(vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _rows_for(s, days, p, n, st):
    return [engine.day_row(s, d, p.get(f"{s['m']}|{s['u']}|{d}", []), n.get(f"{s['m']}|{s['u']}|{d}"), st) for d in days]


def month_sheet(m, days, staff, p, n, st, ms, company="RapidEx"):
    wb = Workbook()
    ws = wb.active
    ws.title = datetime.strptime(m, "%Y-%m").strftime("%b-%Y")
    month_name = datetime.strptime(m, "%Y-%m").strftime("%B-%Y")
    first_day_col, last_day_col = 6, 5 + len(days)
    tot_codes = ["P", "L", "H", "E", "S", "WH"]
    last_col = last_day_col + len(tot_codes)

    # title row
    ws.cell(row=1, column=3, value="Month Start Date").fill = _fill(PLUM)
    ws.cell(row=1, column=3).font = Font(bold=True, color="FF8A3D")
    ws.cell(row=1, column=4, value=datetime.strptime(days[0], "%Y-%m-%d").strftime("%-d-%b-%Y")).fill = _fill(ORANGE)
    ws.cell(row=1, column=4).font = Font(bold=True, color="FFFFFF")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.merge_cells(start_row=1, start_column=first_day_col, end_row=1, end_column=last_day_col)
    t = ws.cell(row=1, column=first_day_col, value=f"{company} Attendance Sheet for {month_name}")
    t.fill, t.font, t.alignment = _fill(PLUM), Font(bold=True, size=16, color="FF8A3D"), CENTER
    ws.row_dimensions[1].height = 30

    # day-name row + header row
    for j, d in enumerate(days):
        c = ws.cell(row=3, column=first_day_col + j, value=engine.WD[datetime.strptime(d, "%Y-%m-%d").weekday()])
        c.fill, c.font, c.alignment = _fill(ORANGE), Font(bold=True, color="FFFFFF"), CENTER
    hdr = ["NO.", "Office", "Department", "Emp Name", "Total Absence"] + \
          [datetime.strptime(d, "%Y-%m-%d").strftime("%d-%b") for d in days] + tot_codes
    _head(ws, 4, hdr)
    for j in range(len(days)):
        ws.cell(row=4, column=first_day_col + j).fill = _fill(ORANGE)
    ws.row_dimensions[4].height = 30

    # body
    r0 = 5
    for i, s in enumerate(staff):
        r = r0 + i
        rows = _rows_for(s, days, p, n, st)
        vals = [i + 1, ms.get(s["m"], ""), s["dep"], s["name"]]
        for c_i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=c_i, value=v)
            c.border = BOX
        rng = f"{get_column_letter(first_day_col)}{r}:{get_column_letter(last_day_col)}{r}"
        tc = ws.cell(row=r, column=5, value=f'=IF(COUNTIF({rng},"A")=0,"-",COUNTIF({rng},"A"))')
        tc.alignment, tc.border, tc.font = CENTER, BOX, Font(bold=True, color="C62828")
        for j, row in enumerate(rows):
            c = ws.cell(row=r, column=first_day_col + j, value=(row["st"] if row else None))
            c.alignment, c.border = CENTER, BOX
            if row and (row["comment"] or row["fixed"]):
                from openpyxl.comments import Comment
                txt = " · ".join(x for x in [row["cin"] and f"In {row['cin']}", row["cout"] and f"Out {row['cout']}", row["comment"]] if x)
                c.comment = Comment(txt, "RapidEx portal")
        for k, code in enumerate(tot_codes):
            c = ws.cell(row=r, column=last_day_col + 1 + k, value=f'=COUNTIF({rng},"{code}")')
            c.alignment, c.border = CENTER, BOX
    last_row = r0 + max(len(staff), 1) - 1

    # code dropdown + colours that follow the code
    if staff:
        area = f"{get_column_letter(first_day_col)}{r0}:{get_column_letter(last_day_col)}{last_row}"
        dv = DataValidation(type="list", formula1='"' + ",".join(engine.CODES) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "Use one of: " + ", ".join(engine.CODES), "Attendance code"
        ws.add_data_validation(dv)
        dv.add(area)
        tl = f"{get_column_letter(first_day_col)}{r0}"
        for code, (bg, fg) in CODE_STYLE.items():
            ws.conditional_formatting.add(area, FormulaRule(formula=[f'{tl}="{code}"'], fill=_fill(bg),
                                                            font=Font(bold=True, color=fg)))
        # highlight people with many absences, like the HR sheet
        ws.conditional_formatting.add(f"D{r0}:D{last_row}", FormulaRule(formula=[f'COUNTIF(${get_column_letter(first_day_col)}{r0}:${get_column_letter(last_day_col)}{r0},"A")>=10'], fill=_fill("FFF59D")))

    _widths(ws, [5, 13, 18, 26, 9] + [7] * len(days) + [5] * len(tot_codes))
    ws.freeze_panes = ws.cell(row=5, column=first_day_col)
    ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_row}"

    # legend + details sheet
    lg = wb.create_sheet("Codes")
    _head(lg, 1, ["Code", "Meaning"])
    for i, (k, v) in enumerate(engine.CODES.items(), 2):
        a = lg.cell(row=i, column=1, value=k)
        bg, fg = CODE_STYLE[k]
        a.fill, a.font, a.alignment = _fill(bg), Font(bold=True, color=fg), CENTER
        lg.cell(row=i, column=2, value=v)
    _widths(lg, [8, 26])

    dt = wb.create_sheet("Daily details")
    _head(dt, 1, ["Date", "Day", "Office", "Machine ID", "Name", "Department", "Code", "Clock In", "Clock Out",
                  "Worked", "Late", "Left Early", "OT", "Comment"])
    out = []
    for s in staff:
        for row in _rows_for(s, days, p, n, st):
            if row:
                out.append([row["ds"], row["day"], ms.get(s["m"], ""), s["u"], s["name"], s["dep"], row["st"], row["cin"],
                            row["cout"], engine.fmt(row["worked"]), engine.fmt(row["late"]) if row["late"] else "",
                            engine.fmt(row["early"]) if row["early"] else "", engine.fmt(row["ot"]) if row["ot"] else "", row["comment"]])
    _rows(dt, 2, sorted(out, key=lambda x: (x[0], x[2], x[4])))
    _widths(dt, [11, 6, 13, 10, 24, 18, 7, 9, 9, 8, 7, 9, 7, 40])
    dt.freeze_panes = "A2"
    return wb


def day_sheet(d, staff, p, n, st, ms):
    wb = Workbook()
    ws = wb.active
    ws.title = d
    _head(ws, 1, ["Office", "Machine ID", "Name", "Designation", "Department", "Shift", "Clock In", "Clock Out",
                  "Worked", "Late", "OT", "Code", "Status", "Comment"])
    rows = []
    for s in staff:
        r = engine.day_row(s, d, p.get(f"{s['m']}|{s['u']}|{d}", []), n.get(f"{s['m']}|{s['u']}|{d}"), st)
        if r:
            rows.append([ms.get(s["m"], ""), s["u"], s["name"], s["des"], s["dep"], f"{r['shift']['in']}-{r['shift']['out']}",
                         r["cin"], r["cout"], engine.fmt(r["worked"]), engine.fmt(r["late"]) if r["late"] else "",
                         engine.fmt(r["ot"]) if r["ot"] else "", r["st"], engine.CODES.get(r["st"], "Not in yet"), r["comment"]])
    _rows(ws, 2, rows)
    _widths(ws, [13, 10, 24, 18, 16, 12, 9, 9, 8, 7, 7, 6, 18, 40])
    ws.freeze_panes = "A2"
    return wb


def report_sheet(s, m, days, p, n, st, ms):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    rows = _rows_for(s, days, p, n, st)
    t = engine.tally(rows)
    month_name = datetime.strptime(m, "%Y-%m").strftime("%B %Y")
    ws.merge_cells("A1:I1")
    c = ws.cell(row=1, column=1, value=f"{s['name']} · Attendance {month_name}")
    c.fill, c.font, c.alignment = _fill(PLUM), Font(bold=True, size=14, color="FFFFFF"), Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    sh = engine.shift_of(s, st)
    info = [("Office", ms.get(s["m"], "")), ("Machine ID", s["u"]), ("Designation", s["des"]), ("Department", s["dep"]),
            ("Timing", f"{sh['in']}-{sh['out']}"), ("Attendance", f"{t['pct']}%"), ("Present", t["P"] + t["L"]),
            ("Late days", t["L"]), ("Absent", t["A"]), ("Leave / WFH", t["S"] + t["E"] + t["H"] + t["WH"]),
            ("Late total", engine.fmt(t["late"])), ("Hours worked", engine.fmt(t["hours"])), ("Overtime", engine.fmt(t["ot"]) if s["ot"] else "Not eligible")]
    for i, (k, v) in enumerate(info):
        r, col = 3 + i // 3, 1 + (i % 3) * 3
        a = ws.cell(row=r, column=col, value=k)
        a.font = Font(bold=True, color="867A8A")
        ws.cell(row=r, column=col + 1, value=v).font = Font(bold=True)
    start = 3 + (len(info) + 2) // 3 + 1
    _head(ws, start, ["Date", "Day", "Code", "Status", "Clock In", "Clock Out", "Worked", "Late", "OT", "Comment"])
    for i, r in enumerate([x for x in rows if x], start + 1):
        vals = [r["ds"], r["day"], r["st"], engine.CODES.get(r["st"], ""), r["cin"], r["cout"], engine.fmt(r["worked"]),
                engine.fmt(r["late"]) if r["late"] else "", engine.fmt(r["ot"]) if r["ot"] else "", r["comment"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BOX
        if r["st"] in CODE_STYLE:
            bg, fg = CODE_STYLE[r["st"]]
            ws.cell(row=i, column=3).fill = _fill(bg)
            ws.cell(row=i, column=3).font = Font(bold=True, color=fg)
            ws.cell(row=i, column=3).alignment = CENTER
    _widths(ws, [12, 13, 14, 18, 12, 14, 9, 8, 8, 40])
    return wb


def ot_sheet(m, days, staff, p, n, st, ms):
    wb = Workbook()
    ws = wb.active
    ws.title = "OT Summary"
    _head(ws, 1, ["Office", "Machine ID", "Name", "Designation", "OT Days", "Total OT (h:mm)", "OT Hours"])
    detail = []
    summ = []
    for s in staff:
        rows = [r for r in _rows_for(s, days, p, n, st) if r and r["ot"]]
        tot = sum(r["ot"] for r in rows)
        summ.append([ms.get(s["m"], ""), s["u"], s["name"], s["des"], len(rows), engine.fmt(tot), round(tot / 60, 2)])
        for r in rows:
            detail.append([r["ds"], r["day"], ms.get(s["m"], ""), s["u"], s["name"], r["cin"], r["cout"],
                           "Off day" if engine.is_off(r["ds"], st) else r["shift"]["out"], engine.fmt(r["ot"]), round(r["ot"] / 60, 2), r["comment"]])
    _rows(ws, 2, sorted(summ, key=lambda x: -x[6]))
    _widths(ws, [13, 10, 24, 18, 9, 15, 10])
    d = wb.create_sheet("OT by day")
    _head(d, 1, ["Date", "Day", "Office", "Machine ID", "Name", "Clock In", "Clock Out", "Office Out", "OT (h:mm)", "OT Hours", "Comment"])
    _rows(d, 2, sorted(detail))
    _widths(d, [11, 6, 13, 10, 24, 9, 9, 10, 10, 9, 36])
    return wb


def punch_sheet(d, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = d
    _head(ws, 1, ["Time", "Office", "Machine ID", "Name", "Tracked"])
    _rows(ws, 2, [[r["ts"][11:], r["office"], r["user_id"], r["name"] or "", "Yes" if r["tracked"] else "No"] for r in rows])
    _widths(ws, [10, 14, 11, 26, 9])
    return wb
